"""Build the controlled, entirely synthetic Word preservation fixture."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

destination = Path(__file__).with_name("board-notes.docx")
if destination.exists():
    raise SystemExit(f"Refusing to overwrite {destination}")
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.85)
for name in ["Normal", "Title", "Heading 1", "List Bullet"]:
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
doc.styles["Title"].font.size = Pt(23)
for border in doc.styles["Title"].element.xpath(".//w:pBdr"):
    border.getparent().remove(border)
doc.styles["Heading 1"].font.size = Pt(14)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
doc.add_paragraph("Facilities committee roof repair notes", "Title")
doc.add_paragraph("The committee deferred roof repair approval pending two additional quotes. These notes record the estimates and the action due before a decision.")
doc.add_paragraph("Meeting notes", "Heading 1")
doc.add_paragraph("Facilities committee, 14 Sep 2026. Roof repair estimate $24,000; last estimate $21,500. Approval deferred. Need two more quotes before a decision. Morgan will request quotes by 21 Sep 2026. No spending authorized today. Attendance not recorded in these notes.")
doc.add_paragraph("Required follow up", "Heading 1")
doc.add_paragraph("Request two additional roof repair quotes.", "List Bullet")
doc.add_paragraph("Return for a decision after the additional quotes are available.", "List Bullet")
doc.add_paragraph("Recorded estimates", "Heading 1")
table = doc.add_table(rows=3, cols=2)
table.autofit = False
table.columns[0].width, table.columns[1].width = Inches(4.6), Inches(2.2)
for row, values in zip(table.rows, [("Estimate", "Amount"), ("Current roof repair estimate", "$24,000"), ("Previous roof repair estimate", "$21,500")]):
    for cell, value in zip(row.cells, values):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].paragraph_format.space_before = Pt(5)
        cell.paragraphs[0].paragraph_format.space_after = Pt(5)
        tcpr = cell._tc.get_or_add_tcPr()
        margins = OxmlElement("w:tcMar")
        for side in ["top", "left", "bottom", "right"]:
            child = OxmlElement("w:" + side)
            child.set(qn("w:w"), "100")
            child.set(qn("w:type"), "dxa")
            margins.append(child)
        tcpr.append(margins)
    row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
for cell in table.rows[0].cells:
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "E9F0F6")
    cell._tc.get_or_add_tcPr().append(shade)
    cell.paragraphs[0].runs[0].bold = True
borders = OxmlElement("w:tblBorders")
for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
    border = OxmlElement("w:" + side)
    for attr, value in [("val", "single"), ("sz", "6"), ("color", "D9D9D9")]:
        border.set(qn("w:" + attr), value)
    borders.append(border)
table._tbl.tblPr.append(borders)
doc.add_paragraph("Record keeping", "Heading 1")
p = doc.add_paragraph()
p.add_run("Spending restriction. ").bold = True
p.add_run("No spending was authorized today. ")
p.add_run("Attendance was not recorded.").italic = True
doc.core_properties.title = "Facilities committee roof repair notes"
doc.core_properties.subject = "Synthetic local Word adapter verification fixture"
doc.core_properties.author = "Local Gemma verification"
# Keep the synthetic fixture independent of a local build date.
package = BytesIO()
doc.save(package)
with ZipFile(package) as archive:
    members = [(info, archive.read(info)) for info in archive.infolist()]
with ZipFile(destination, "w") as archive:
    for info, contents in members:
        info.date_time = (1980, 1, 1, 0, 0, 0)
        archive.writestr(info, contents)
print(destination)
