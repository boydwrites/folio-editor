"""Local-only text import and clean Word exports for the editor.

Imports never follow document links or execute embedded content. Template
imports copy a small set of formatting values into a fresh document, never
the source package, its text, images, headers, comments, or metadata.
"""

from __future__ import annotations

import base64
import binascii
from datetime import date as Date
from io import BytesIO
import math
from pathlib import PurePosixPath
import re
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_TEXT_CHARS = 500_000
MAX_ZIP_BYTES = 40 * 1024 * 1024
MAX_XML_BYTES = 8 * 1024 * 1024
MAX_PDF_PAGES = 100
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
FONTS = {"Aptos", "Calibri", "Georgia", "Times New Roman"}
STYLE_NAMES = ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number")


def _checked_text(text: str) -> str:
    if not isinstance(text, str):
        raise ValueError("Provide your notes as text.")
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("This document is too long. Import up to 500,000 characters at a time.")
    if not text:
        raise ValueError("This file contains no readable text. Paste the text, or import a Word, text, or text-based PDF file.")
    if any((ord(char) < 32 and char not in "\n\t") or 0xD800 <= ord(char) <= 0xDFFF or ord(char) in (0xFFFE, 0xFFFF) for char in text):
        raise ValueError("This file contains unsupported control characters. Save it as UTF-8 text or DOCX.")
    return text


def _open_docx(data: bytes) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(BytesIO(data))
        members = archive.infolist()
        if len(members) > 1000 or sum(m.file_size for m in members) > MAX_ZIP_BYTES:
            raise ValueError("This Word document expands beyond the import limit. Save a smaller copy.")
        seen = set()
        for member in members:
            name = member.filename
            if name in seen or "\\" in name or PurePosixPath(name).is_absolute() or ".." in PurePosixPath(name).parts:
                raise ValueError("This Word document has an invalid package structure. Save a new DOCX copy.")
            seen.add(name)
            if member.flag_bits & 1 or member.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                raise ValueError("This Word document uses unsupported encryption or compression.")
            if member.file_size > 1024 * 1024 and member.file_size / max(1, member.compress_size) > 250:
                raise ValueError("This Word document expands beyond the import limit. Save a smaller copy.")
            if name.endswith((".xml", ".rels")) and member.file_size > MAX_XML_BYTES:
                raise ValueError("This Word document is too large to import. Save a smaller copy.")
        if "word/document.xml" not in seen:
            raise ValueError("This file is not a valid Word document. Save it as DOCX and try again.")
        return archive
    except ValueError:
        if "archive" in locals():
            archive.close()
        raise
    except (zipfile.BadZipFile, OSError, RuntimeError):
        raise ValueError("This Word document could not be read. Save a new DOCX copy and try again.") from None


def _xml(archive: zipfile.ZipFile, name: str):
    try:
        data = archive.read(name)
        if len(data) > MAX_XML_BYTES:
            raise ValueError("This Word document is too large to import.")
        parser = etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True, huge_tree=False)
        root = etree.fromstring(data, parser=parser)
        if root.getroottree().docinfo.doctype or any(isinstance(node, etree._Entity) for node in root.iter()):
            raise ValueError("This Word document contains unsupported XML entities. Save a new DOCX copy.")
        return root
    except (etree.XMLSyntaxError, zipfile.BadZipFile, KeyError, RuntimeError):
        raise ValueError("This Word document has invalid content. Save a new DOCX copy and try again.") from None


def _paragraph_text(element) -> str:
    pieces = []

    def visit(node):
        if node.tag in (W + "del", W + "moveFrom", W + "pPr", W + "rPr"):
            return
        if node.tag == W + "r":
            properties = node.find(W + "rPr")
            if properties is not None:
                for kind in ("vanish", "webHidden"):
                    value = properties.find(W + kind)
                    if value is not None and value.get(W + "val", "true") not in ("0", "false", "off"):
                        return
        if node.tag == W + "t":
            pieces.append(node.text or "")
        elif node.tag == W + "tab":
            pieces.append("\t")
        elif node.tag in (W + "br", W + "cr"):
            pieces.append("\n")
        else:
            for child in node:
                visit(child)

    visit(element)
    return "".join(pieces)


def _block_text(element) -> list[str]:
    if element.tag == W + "p":
        return [_paragraph_text(element)]
    if element.tag == W + "tbl":
        rows = []
        for row in element.findall(W + "tr"):
            cells = []
            for cell in row.findall(W + "tc"):
                cells.append("\n".join(text for child in cell for text in _block_text(child)).strip())
            rows.append("\t".join(cells))
        return rows
    if element.tag in (W + "body", W + "sdt", W + "sdtContent", W + "customXml", W + "ins"):
        return [text for child in element for text in _block_text(child)]
    return []


def extract_text(name: str, data: bytes) -> str:
    """Extract the visible body locally; raise ValueError for unsupported input."""
    if not isinstance(data, bytes) or not data:
        raise ValueError("Choose a file that contains your notes.")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("This file is too large. Choose a file smaller than 15 MB.")
    extension = PurePosixPath(name).suffix.lower()
    if extension in (".txt", ".md"):
        try:
            encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            return _checked_text(data.decode(encoding))
        except UnicodeError:
            raise ValueError("This text file is not UTF-8 or UTF-16. Save it as UTF-8, or paste its contents.") from None
    if extension == ".doc":
        raise ValueError("Older .doc files are not supported. Save this file as .docx in Word and import that copy.")
    if extension == ".docx":
        with _open_docx(data) as archive:
            root = _xml(archive, "word/document.xml")
            if root.tag != W + "document":
                raise ValueError("This file is not a valid Word document.")
            body = root.find(W + "body")
            return _checked_text("\n\n".join(_block_text(body)) if body is not None else "")
    if extension == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ValueError("This PDF is password-protected. Import an unlocked copy, or paste the text.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ValueError("This PDF is too long. Import up to 100 pages at a time.")
            pages = []
            size = 0
            for page in reader.pages:
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > MAX_XML_BYTES:
                    raise ValueError("This PDF page is too complex to import. Save a text-only copy.")
                value = page.extract_text() or ""
                if not value.strip():
                    raise ValueError("This PDF has a page with no selectable text, which usually means it was scanned. The editor can't read scanned pages. Paste the text, or import a text-based copy.")
                size += len(value)
                if size > MAX_TEXT_CHARS:
                    raise ValueError("This document is too long. Import up to 500,000 characters at a time.")
                pages.append(value)
            return _checked_text("\n\n".join(pages))
        except ValueError:
            raise
        except ImportError:
            raise ValueError("PDF import is unavailable in this Python environment. Paste the text, or import a Word file.") from None
        except Exception:
            raise ValueError("This PDF could not be read. Save a text-based copy or paste your notes.") from None
    raise ValueError("The editor can import .txt, .md, .docx, or text-based .pdf files.")


def _number(value, default: float, lower: float, upper: float) -> float:
    try:
        value = float(value)
        return min(upper, max(lower, value)) if math.isfinite(value) else default
    except (TypeError, ValueError, OverflowError):
        return default


def _template_values(template: dict) -> dict:
    """Read only safe scalar style/layout values; never copy XML or relationships."""
    if not isinstance(template, dict) or not str(template.get("name", "")).lower().endswith(".docx"):
        raise ValueError("Choose a DOCX file for the document format reference.")
    encoded = template.get("data", "")
    if not isinstance(encoded, str) or len(encoded) > (MAX_FILE_BYTES * 4 // 3 + 4):
        raise ValueError("The document format reference is too large.")
    try:
        data = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError):
        raise ValueError("The document format reference could not be read. Import it again.") from None
    values = {"styles": {}}
    with _open_docx(data) as archive:
        root = _xml(archive, "word/document.xml")
        body = root.find(W + "body")
        section = body.find(W + "sectPr") if body is not None else None
        if section is not None:
            size, margin = section.find(W + "pgSz"), section.find(W + "pgMar")
            if size is not None:
                values["width"] = _number(size.get(W + "w"), 12240, 8640, 24480) / 1440
                values["height"] = _number(size.get(W + "h"), 15840, 8640, 24480) / 1440
            if margin is not None:
                values["margins"] = {edge: _number(margin.get(W + edge), 1440, 720, 2160) / 1440 for edge in ("top", "right", "bottom", "left")}
        if "word/styles.xml" in archive.namelist():
            styles = _xml(archive, "word/styles.xml")
            for style in styles.findall(W + "style"):
                identifier = style.get(W + "styleId", "")
                target = next((name for name in STYLE_NAMES if name.replace(" ", "") == identifier), None)
                if target is None:
                    continue
                result = {}
                run, paragraph = style.find(W + "rPr"), style.find(W + "pPr")
                if run is not None:
                    font, size = run.find(W + "rFonts"), run.find(W + "sz")
                    if font is not None and font.get(W + "ascii") in FONTS:
                        result["font"] = font.get(W + "ascii")
                    if size is not None:
                        result["size"] = _number(size.get(W + "val"), 22, 20, 56) / 2
                    for prop in ("b", "i"):
                        node = run.find(W + prop)
                        if node is not None:
                            result[prop] = node.get(W + "val", "true") not in ("0", "false", "off")
                spacing = paragraph.find(W + "spacing") if paragraph is not None else None
                if spacing is not None:
                    for prop in ("before", "after"):
                        if spacing.get(W + prop) is not None:
                            result[prop] = _number(spacing.get(W + prop), 0, 0, 480) / 20
                    if spacing.get(W + "lineRule", "auto") == "auto" and spacing.get(W + "line"):
                        result["lineSpacing"] = _number(spacing.get(W + "line"), 276, 240, 384) / 240
                values["styles"][target] = result
    return values


def _set_font(style, name: str, size: float) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    properties = style.element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is not None:
        for attribute in tuple(fonts.attrib):
            if attribute.endswith("Theme"):
                del fonts.attrib[attribute]
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attribute), name)


def _emphasis_close(text: str, start: int, count: int) -> int:
    """Match the editor's emphasis subset, including nested bold and italic."""
    character = text[start]
    nested = []
    cursor = start + count
    while cursor < len(text):
        if text[cursor] == "\\":
            cursor += 2
            continue
        if text[cursor] != character:
            cursor += 1
            continue
        end = cursor
        while end < len(text) and text[end] == character:
            end += 1
        remaining, at = end - cursor, cursor
        can_close = not text[cursor - 1].isspace() and not (
            character == "_" and end < len(text) and text[end].isalnum()
        )
        can_open = end < len(text) and not text[end].isspace() and not (
            character == "_" and cursor > 0 and text[cursor - 1].isalnum()
        )
        if can_close:
            while nested and remaining >= nested[-1]:
                consumed = nested.pop()
                at += consumed
                remaining -= consumed
            if not nested and remaining >= count:
                return at
        if can_open and remaining:
            nested.append(min(3, remaining))
        cursor = end
    return -1


def _inline(paragraph, text: str) -> None:
    # Deliberately small Markdown subset: never interpret HTML, links or images.
    # Escapes are decoded only after block detection, so a literal "2026\."
    # stays a paragraph and never becomes an ordered-list marker.
    segments = []

    def append(value, bold, italic):
        if segments and segments[-1][1:] == (bold, italic):
            segments[-1] = (segments[-1][0] + value, bold, italic)
        else:
            segments.append((value, bold, italic))

    def parse(value, bold=False, italic=False, depth=0):
        index = 0
        while index < len(value):
            character = value[index]
            if character == "\\" and index + 1 < len(value) and value[index + 1] in "\\*_[]#!+-.>()`":
                append(value[index + 1], bold, italic)
                index += 2
                continue
            if depth < 32 and character in "*_" and not (character == "_" and index > 0 and value[index - 1].isalnum()):
                count = 1
                while count < 3 and index + count < len(value) and value[index + count] == character:
                    count += 1
                closing = _emphasis_close(value, index, count) if index + count < len(value) and not value[index + count].isspace() else -1
                if closing > index + count:
                    parse(value[index + count:closing], bold or count >= 2, italic or count in (1, 3), depth + 1)
                    index = closing + count
                    continue
            append(character, bold, italic)
            index += 1

    parse(text)
    for segment, bold, italic in segments:
        run = paragraph.add_run(segment)
        if bold:
            run.bold = True
        if italic:
            run.italic = True


def export_docx(title: str, text: str, settings: dict, date: str | None = None) -> bytes:
    """Create a clean DOCX with the supplied date and body text.

    A leading Markdown H1 supplies the displayed Title; otherwise title does.
    The title argument always supplies document metadata, even when a leading
    H1 has a different wording. Later headings remain in their body positions.

    settings.template may be {name: "format.docx", data: "<base64>"}. Its known
    fonts, type sizes, paragraph spacing, page dimensions and margins are read
    as a format reference. Explicit settings take precedence over that reference.
    """
    text = _checked_text(text)
    title = _checked_text(title) if title and title.strip() else ""
    if len(title) > 300 or "\n" in title:
        raise ValueError("Use a document title of up to 300 characters on one line.")
    lines = text.splitlines()
    leading_title = re.match(r"^ {0,3}#\s+(.+?)(?:\s+#+)?\s*$", lines[0])
    # A leading lower-level heading that repeats the title would print twice.
    leading_repeat = re.match(r"^ {0,3}#{2,3}\s+(.+?)(?:\s+#+)?\s*$", lines[0])
    same = lambda a, b: " ".join(a.split()).casefold() == " ".join(b.split()).casefold()
    if not leading_title and leading_repeat and title and same(leading_repeat.group(1), title):
        leading_title = leading_repeat
    displayed_title = leading_title.group(1) if leading_title else title
    if leading_title:
        lines = lines[1:]
    settings = settings if isinstance(settings, dict) else {}
    template = _template_values(settings["template"]) if settings.get("template") else {}
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(template.get("width", 8.5))
    section.page_height = Inches(template.get("height", 11))
    for edge in ("top", "right", "bottom", "left"):
        setattr(section, edge + "_margin", Inches(template.get("margins", {}).get(edge, 1)))
    if "pageSize" in settings:
        dimensions = (8.2677, 11.6929) if settings["pageSize"] == "a4" else (8.5, 11)
        section.page_width, section.page_height = (Inches(value) for value in dimensions)
    if "margins" in settings:
        for edge in ("top", "right", "bottom", "left"):
            setattr(section, edge + "_margin", Inches(.75 if settings["margins"] == "compact" else 1))
    section.header_distance = Inches(.4)
    section.footer_distance = Inches(.4)
    base_size = _number(settings.get("fontSize"), 11, 10, 14)
    for name in STYLE_NAMES:
        style = document.styles[name]
        reference = template.get("styles", {}).get(name, {})
        font = settings.get("font", reference.get("font", "Aptos"))
        font = font if isinstance(font, str) and font in FONTS else "Aptos"
        default_size = {"Title": 23, "Heading 1": 14, "Heading 2": 12, "Heading 3": 11.5}.get(name, base_size)
        size = reference.get("size", default_size)
        if "fontSize" in settings:
            size = {"Title": base_size + 12, "Heading 1": base_size + 3, "Heading 2": base_size + 1, "Heading 3": base_size + .5}.get(name, base_size)
        _set_font(style, font, size)
        style.font.bold = reference.get("b", name == "Title" or name.startswith("Heading"))
        style.font.italic = reference.get("i", False)
        style.font.underline = False
        properties = style.element.find(qn("w:pPr"))
        if properties is not None:
            for tag in ("w:pBdr", "w:shd"):
                for decoration in properties.findall(qn(tag)):
                    properties.remove(decoration)
        paragraph = style.paragraph_format
        paragraph.line_spacing = _number(settings.get("lineSpacing", reference.get("lineSpacing")), 1.15, 1, 1.6)
        paragraph.space_before = Pt(reference.get("before", 15 if name.startswith("Heading") else 0))
        paragraph.space_after = Pt(reference.get("after", 14 if name == "Title" else 7))
        paragraph.widow_control = True
        paragraph.keep_with_next = name == "Title" or name.startswith("Heading")
    if displayed_title:
        paragraph = document.add_paragraph(style="Title")
        _inline(paragraph, displayed_title)
    if settings.get("includeDate") and date:
        try:
            parsed = Date.fromisoformat(date)
            date_text = f"{parsed.strftime('%B')} {parsed.day}, {parsed.year}"
        except (TypeError, ValueError):
            date_text = _checked_text(str(date))
        paragraph = document.add_paragraph(date_text)
        paragraph.paragraph_format.space_after = Pt(14)
        paragraph.paragraph_format.keep_with_next = True
    for line in lines:
        if not line.strip():
            continue
        heading = re.match(r"^ {0,3}(#{1,3})\s+(.+?)(?:\s+#+)?\s*$", line)
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        number = re.match(r"^\s*(\d+[.)])\s+(.+)$", line)
        if heading:
            paragraph = document.add_paragraph(style=f"Heading {len(heading.group(1))}")
            _inline(paragraph, heading.group(2))
        elif bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _inline(paragraph, bullet.group(1))
        elif number:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(.25)
            paragraph.paragraph_format.first_line_indent = Inches(-.25)
            _inline(paragraph, number.group(1) + "  " + number.group(2))
        else:
            paragraph = document.add_paragraph()
            _inline(paragraph, line.strip())
    if settings.get("confidentiality"):
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("Confidential")
        run.font.name = document.styles["Normal"].font.name
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)
    # Fresh package + blank personal properties prevent source content leakage.
    properties = document.core_properties
    for field in ("author", "last_modified_by", "comments", "keywords", "subject", "category", "identifier", "language", "version", "content_status"):
        setattr(properties, field, "")
    properties.title = title
    for tag in ("dcterms:created", "dcterms:modified", "cp:lastPrinted"):
        node = properties._element.find(qn(tag))
        if node is not None:
            properties._element.remove(node)
    properties.revision = 1
    result = BytesIO()
    document.save(result)
    return result.getvalue()
