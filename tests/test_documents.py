"""Import safety and content/layout preservation tests using synthetic notes."""

import base64
from io import BytesIO
from pathlib import Path
import sys
import unittest
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.shared import Inches, Pt
from lxml import etree
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from folio.documents import extract_text, export_docx, MAX_FILE_BYTES, W


def package(document_xml, extra=None, compression=zipfile.ZIP_DEFLATED):
    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression) as archive:
        archive.writestr("word/document.xml", document_xml)
        for name, content in (extra or {}).items():
            archive.writestr(name, content)
    return output.getvalue()


def wrapped(body):
    return f'<w:document xmlns:w="{W[1:-1]}"><w:body>{body}</w:body></w:document>'


class ImportTests(unittest.TestCase):
    def test_unicode_text_and_line_endings(self):
        self.assertEqual(extract_text("notes.TXT", "\ufeffBudget — €40\r\nApproved".encode()), "Budget — €40\nApproved")
        self.assertEqual(extract_text("notes.md", "Board notes".encode("utf-16")), "Board notes")

    def test_unsupported_legacy_word_and_binary_text(self):
        for name, data, match in [("notes.doc", b"data", "Save this file as .docx"), ("notes.exe", b"data", "Folio can import"), ("notes.txt", b"null\x00", "control characters")]:
            with self.subTest(name=name), self.assertRaisesRegex(ValueError, match):
                extract_text(name, data)

    def test_large_upload_rejected_before_parse(self):
        with self.assertRaisesRegex(ValueError, "15 MB"):
            extract_text("notes.docx", b"x" * (MAX_FILE_BYTES + 1))

    def test_body_table_and_paragraph_order(self):
        document = Document()
        document.add_paragraph("Opening decision")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Owner"
        table.cell(0, 1).text = "Action"
        table.cell(1, 0).text = "Maya"
        table.cell(1, 1).text = "Confirm vendor\nReport Friday"
        document.add_paragraph("Next meeting")
        output = BytesIO()
        document.save(output)
        self.assertEqual(extract_text("notes.docx", output.getvalue()), "Opening decision\n\nOwner\tAction\n\nMaya\tConfirm vendor\nReport Friday\n\nNext meeting")

    def test_deleted_and_hidden_text_are_excluded(self):
        xml = wrapped('<w:p><w:r><w:t>Approved </w:t></w:r><w:del><w:r><w:delText>rejected</w:delText></w:r></w:del><w:r><w:rPr><w:vanish/></w:rPr><w:t>hidden secret</w:t></w:r><w:ins><w:r><w:t>budget</w:t></w:r></w:ins></w:p>')
        self.assertEqual(extract_text("notes.docx", package(xml)), "Approved budget")

    def test_formatting_tab_stops_are_not_imported_as_text(self):
        xml = wrapped('<w:p><w:pPr><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs></w:pPr><w:r><w:t>Owner</w:t><w:tab/><w:t>Maya</w:t></w:r></w:p>')
        self.assertEqual(extract_text("notes.docx", package(xml)), "Owner\tMaya")

    def test_xml_external_entities_rejected(self):
        xml = f'<!DOCTYPE doc [<!ENTITY secret SYSTEM "file:///etc/passwd">]>{wrapped("<w:p><w:r><w:t>&secret;</w:t></w:r></w:p>")}'
        with self.assertRaisesRegex(ValueError, "XML entities"):
            extract_text("notes.docx", package(xml))

    def test_zip_path_traversal_and_expansion_rejected(self):
        with self.assertRaisesRegex(ValueError, "package structure"):
            extract_text("notes.docx", package(wrapped(""), {"../outside": "bad"}))
        with self.assertRaisesRegex(ValueError, "expands beyond"):
            extract_text("notes.docx", package(wrapped(""), {"word/large.xml": "A" * (2 * 1024 * 1024)}))

    def test_duplicate_zip_members_rejected(self):
        import warnings
        output = BytesIO()
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            with zipfile.ZipFile(output, "w") as archive:
                archive.writestr("word/document.xml", wrapped(""))
                archive.writestr("word/document.xml", wrapped(""))
        with self.assertRaisesRegex(ValueError, "package structure"):
            extract_text("notes.docx", output.getvalue())

    def test_pdf_extracts_selectable_text(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 720 Td (Budget approved.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        output = BytesIO()
        writer.write(output)
        self.assertIn("Budget approved.", extract_text("notes.pdf", output.getvalue()))

    def test_pdf_scans_and_encryption_get_actionable_errors(self):
        for encrypted, match in [(False, "can.t read scanned pages"), (True, "password-protected")]:
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            if encrypted:
                writer.encrypt("test-password")
            output = BytesIO()
            writer.write(output)
            with self.subTest(encrypted=encrypted), self.assertRaisesRegex(ValueError, match):
                extract_text("notes.pdf", output.getvalue())


class ExportTests(unittest.TestCase):
    def test_rich_editor_bold_italic_and_nested_emphasis_export_as_native_runs(self):
        text = "## Finance\n*The **board** approved* **the *revised* budget** and ***next steps***.\n__María__ confirmed _€40,000_ for October 19."
        document = Document(BytesIO(export_docx("Board Meeting Notes", text, {})))
        paragraph = document.paragraphs[2]
        self.assertEqual(paragraph.text, "The board approved the revised budget and next steps.")
        styles = {run.text.strip(): (bool(run.bold), bool(run.italic)) for run in paragraph.runs if run.text.strip()}
        self.assertEqual(styles["The"], (False, True))
        self.assertEqual(styles["board"], (True, True))
        self.assertEqual(styles["revised"], (True, True))
        self.assertEqual(styles["budget"], (True, False))
        self.assertEqual(styles["next steps"], (True, True))
        self.assertEqual(document.paragraphs[3].text, "María confirmed €40,000 for October 19.")
        self.assertTrue(next(run for run in document.paragraphs[3].runs if run.text == "€40,000").italic)

    def test_rich_editor_escaped_literals_preserve_text_and_do_not_become_lists_or_headings(self):
        text = r"2026\. Revenue was $2.4m, up 12.5%." + "\n" + r"\# Decisions remain pending." + "\n" + r"\- This is a literal dash." + "\n" + r"Literal \*stars\* and board\_notes; C:\\Notes\\2026." + "\n" + "board_notes and unmatched * stay unchanged."
        document = Document(BytesIO(export_docx("Board Meeting Notes", text, {})))
        body = document.paragraphs[1:]
        self.assertEqual([paragraph.text for paragraph in body], [
            "2026. Revenue was $2.4m, up 12.5%.",
            "# Decisions remain pending.",
            "- This is a literal dash.",
            r"Literal *stars* and board_notes; C:\Notes\2026.",
            "board_notes and unmatched * stay unchanged.",
        ])
        self.assertTrue(all(paragraph.style.name == "Normal" for paragraph in body))
        self.assertTrue(all(paragraph.paragraph_format.left_indent is None for paragraph in body))
        self.assertFalse(any(run.italic or run.bold for paragraph in body for run in paragraph.runs))

    def test_rich_editor_html_and_links_remain_literal_without_external_relationships(self):
        text = '<img src="https://example.invalid/private">\n[Reference](https://example.invalid/notes)\nA **decision** with *conditions*.'
        data = export_docx("Board Meeting Notes", text, {})
        document = Document(BytesIO(data))
        self.assertEqual(document.paragraphs[1].text, '<img src="https://example.invalid/private">')
        self.assertEqual(document.paragraphs[2].text, '[Reference](https://example.invalid/notes)')
        self.assertEqual(document.paragraphs[3].text, 'A decision with conditions.')
        with zipfile.ZipFile(BytesIO(data)) as archive:
            for name in archive.namelist():
                if name.endswith('.rels'):
                    self.assertNotIn(b'TargetMode="External"', archive.read(name))

    def test_leading_h1_replaces_visible_title_and_keeps_metadata_date_and_facts(self):
        text = "\n\n# September 14 Board Meeting\n\n**Budget review**\n- Approved **$250,000**, subject to legal review.\n\n# Next meeting\nOctober 12."
        data = export_docx("September board meeting", text, {"includeDate": True}, "2026-09-14")
        document = Document(BytesIO(data))
        paragraphs = document.paragraphs
        self.assertEqual(document.core_properties.title, "September board meeting")
        self.assertEqual([p.text for p in paragraphs if p.style.name == "Title"], ["September 14 Board Meeting"])
        self.assertEqual([p.text for p in paragraphs], ["September 14 Board Meeting", "September 14, 2026", "Budget review", "Approved $250,000, subject to legal review.", "Next meeting", "October 12."])
        self.assertEqual(paragraphs[4].style.name, "Heading 1")
        self.assertEqual(paragraphs[3].style.name, "List Bullet")
        self.assertTrue(next(run for run in paragraphs[3].runs if run.text == "$250,000").bold)

    def test_leading_subheading_that_repeats_the_title_prints_once(self):
        paragraphs = Document(BytesIO(export_docx("Project update", "## Project  Update\n\nThe team met.", {}))).paragraphs
        self.assertEqual([(p.style.name, p.text) for p in paragraphs if p.text], [("Title", "Project  Update"), ("Normal", "The team met.")])
        paragraphs = Document(BytesIO(export_docx("Notes", "## Other\n\nThe team met.", {}))).paragraphs
        self.assertEqual([(p.style.name, p.text) for p in paragraphs if p.text], [("Title", "Notes"), ("Heading 2", "Other"), ("Normal", "The team met.")])

    def test_empty_title_prints_no_title(self):
        paragraphs = Document(BytesIO(export_docx("", "The garden was quiet.", {}))).paragraphs
        self.assertEqual([(p.style.name, p.text) for p in paragraphs if p.text], [("Normal", "The garden was quiet.")])

    def test_exports_native_structure_without_changing_facts(self):
        data = export_docx("September Board Notes", "## Decisions\nThe board approved **$250,000** for the pilot.\n\n- Maya will confirm the timeline.\n3. Review the revised forecast.", {}, "2026-09-14")
        document = Document(BytesIO(data))
        paragraphs = document.paragraphs
        self.assertEqual(paragraphs[0].style.name, "Title")
        self.assertEqual(paragraphs[1].style.name, "Heading 2")
        self.assertEqual(paragraphs[3].style.name, "List Bullet")
        self.assertEqual(paragraphs[2].text, "The board approved $250,000 for the pilot.")
        self.assertTrue(next(run for run in paragraphs[2].runs if "$250,000" in run.text).bold)
        self.assertTrue(paragraphs[4].text.startswith("3."))
        self.assertNotIn("September 14", extract_text("export.docx", data))

    def test_font_page_date_footer_and_metadata(self):
        data = export_docx("Board Notes", "Approved subject to legal review.", {"font": "Georgia", "fontSize": 12, "lineSpacing": 1.5, "margins": "compact", "pageSize": "a4", "includeDate": True, "confidentiality": True}, "2026-09-14")
        document = Document(BytesIO(data))
        self.assertEqual(document.styles["Normal"].font.name, "Georgia")
        self.assertEqual(document.styles["Normal"].font.size.pt, 12)
        self.assertEqual(document.styles["Normal"].paragraph_format.line_spacing, 1.5)
        self.assertAlmostEqual(document.sections[0].left_margin.inches, .75, places=3)
        self.assertAlmostEqual(document.sections[0].page_width.inches, 8.2677, places=3)
        self.assertIn("September 14, 2026", [p.text for p in document.paragraphs])
        self.assertEqual(document.sections[0].footer.paragraphs[0].text, "Confidential")
        self.assertEqual(document.core_properties.author, "")
        self.assertEqual(document.core_properties.last_modified_by, "")
        self.assertIsNone(document.core_properties.created)

    def test_invalid_settings_do_not_produce_unreadable_documents(self):
        document = Document(BytesIO(export_docx("Board Notes", "Body.", {"font": "Unknown", "fontSize": 900, "lineSpacing": float("nan")})))
        self.assertEqual(document.styles["Normal"].font.name, "Aptos")
        self.assertEqual(document.styles["Normal"].font.size.pt, 14)
        self.assertAlmostEqual(document.styles["Normal"].paragraph_format.line_spacing, 1.15)

    def test_template_retains_scalar_formatting_without_source_content(self):
        source = Document()
        source.styles["Normal"].font.name = "Georgia"
        source.styles["Normal"].font.size = Pt(12)
        source.styles["Normal"].paragraph_format.line_spacing = 1.5
        source.sections[0].left_margin = Inches(.8)
        source.add_paragraph("SECRET OLD BOARD CONTENT")
        source.sections[0].header.paragraphs[0].text = "SECRET HEADER"
        source.sections[0].footer.paragraphs[0].text = "SECRET FOOTER"
        source.core_properties.author = "SECRET AUTHOR"
        source.core_properties.comments = "SECRET COMMENTS"
        output = BytesIO()
        source.save(output)
        template = {"name": "format.docx", "data": base64.b64encode(output.getvalue()).decode()}
        data = export_docx("New Notes", "The board requested a revised forecast.", {"template": template})
        document = Document(BytesIO(data))
        self.assertEqual(document.styles["Normal"].font.name, "Georgia")
        self.assertEqual(document.styles["Normal"].font.size.pt, 12)
        self.assertEqual(document.styles["Normal"].paragraph_format.line_spacing, 1.5)
        self.assertAlmostEqual(document.sections[0].left_margin.inches, .8, places=3)
        with zipfile.ZipFile(BytesIO(data)) as archive:
            for name in archive.namelist():
                self.assertNotIn(b"SECRET", archive.read(name), name)
            self.assertFalse(any("header" in name or "footer" in name for name in archive.namelist()))
        overridden = Document(BytesIO(export_docx("New", "Body.", {"template": template, "font": "Calibri", "margins": "standard"})))
        self.assertEqual(overridden.styles["Normal"].font.name, "Calibri")
        self.assertEqual(overridden.sections[0].left_margin.inches, 1)

    def test_title_and_heading_styles_have_no_color_theme_or_decorative_rule(self):
        data = export_docx("Board Notes", "# Decisions\nReview Friday.", {})
        with zipfile.ZipFile(BytesIO(data)) as archive:
            styles = etree.fromstring(archive.read("word/styles.xml"))
            for identifier in ("Title", "Heading1", "Heading2", "Heading3"):
                node = next(node for node in styles if node.get(W + "styleId") == identifier)
                self.assertEqual(node.find(W + "rPr").find(W + "color").get(W + "val"), "000000")
                self.assertIsNone(node.find(".//" + W + "pBdr"))


if __name__ == "__main__":
    unittest.main()
